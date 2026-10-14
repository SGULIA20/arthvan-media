import json

import pytest

openpyxl = pytest.importorskip("openpyxl")              # the CI machine does not have it; the Mac does
from openpyxl import load_workbook  # noqa: E402

from engine import review  # noqa: E402


def make(tmp_path):
    root = tmp_path / "repo"; (root / "stories").mkdir(parents=True); (root / "content" / "translations" / "hi").mkdir(parents=True)
    (root / "stories" / "a_en.json").write_text(json.dumps({"id": "a-en", "app": "scoremybiz", "title": "T", "caption": "C", "scenes": [{"vo": "Hello", "card": "Card"}, {"vo": "End"}]}))
    (root / "content" / "translations" / "hi" / "a-en.json").write_text(json.dumps({"title": "शीर्षक", "caption": "कैप्शन", "machine_translated": True, "scenes": [{"vo": "नमस्ते", "card": "कार्ड"}, {"vo": "अंत"}]}))
    return root


def test_export_lists_every_line_card_title_and_caption_beside_the_english(tmp_path):
    root = make(tmp_path)
    p = review.export("hi", root, tmp_path / "Review")
    ws = load_workbook(p).active
    got = [(r[2], r[3], r[4]) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert got == [(1, "Hello", "नमस्ते"), ("1 (on-screen card)", "Card", "कार्ड"), (2, "End", "अंत"), ("title", "T", "शीर्षक"), ("caption", "C", "कैप्शन")]


def test_corrections_typed_in_the_sheet_replace_lines_and_survive_a_re_export(tmp_path):
    root = make(tmp_path); folder = tmp_path / "Review"
    p = review.export("hi", root, folder)
    wb = load_workbook(p); ws = wb.active; ws["F2"] = "  नमस्कार  "; ws["F6"] = "नया कैप्शन"; wb.save(p)
    assert review.apply("hi", root, folder) == {"a-en": 2}
    tr = json.loads((root / "content" / "translations" / "hi" / "a-en.json").read_text())
    assert tr["scenes"][0]["vo"] == "नमस्कार" and tr["caption"] == "नया कैप्शन" and tr["reviewed"] is True and "machine_translated" not in tr
    p2 = review.export("hi", root, folder)                                               # a new export keeps what the reader typed
    assert load_workbook(p2).active["F2"].value == "नमस्कार"
