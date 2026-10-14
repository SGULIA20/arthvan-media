"""Native-speaker review of translated Reel scripts.
python -m engine.review export hi    writes ~/Desktop/Arthvan-Reels/Review/Hindi-review.xlsx : one row per spoken line, English beside the draft,
                                     with a blank 'Correction' column for the reader. Rows already corrected before are kept.
python -m engine.review apply hi     reads the same file back: every non-empty Correction replaces that line in content/translations/<lang>/<id>.json
                                     (and clears the machine_translated flag once a file has been reviewed). Re-render the Reels after applying."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .library import APPS, LANGUAGES

ROOT = Path(__file__).resolve().parent.parent
FOLDER = Path.home() / "Desktop" / "Arthvan-Reels" / "Review"
HEAD = ["App", "Reel", "Line", "English", "Draft", "Correction (write the better version here, leave empty if the draft is fine)", "Key"]


def rows(lang: str, root: Path = ROOT):
    out = []
    for f in sorted((root / "content" / "translations" / lang).glob("*.json")):
        tr = json.loads(f.read_text()); base = json.loads((root / "stories" / (f.stem.replace("-", "_") + ".json")).read_text())
        for i, (b, t) in enumerate(zip(base["scenes"], tr["scenes"])):
            out.append((APPS.get(base["app"], base["app"]), base["id"], i + 1, b.get("vo", ""), t.get("vo", ""), "", f"{f.stem}|{i}|vo"))
            if b.get("card"):
                out.append((APPS.get(base["app"], base["app"]), base["id"], f"{i + 1} (on-screen card)", b["card"], t.get("card", ""), "", f"{f.stem}|{i}|card"))
        out.append((APPS.get(base["app"], base["app"]), base["id"], "title", base.get("title", ""), tr.get("title", ""), "", f"{f.stem}|-|title"))
        out.append((APPS.get(base["app"], base["app"]), base["id"], "caption", base.get("caption", ""), tr.get("caption", ""), "", f"{f.stem}|-|caption"))
    return out


def export(lang: str, root: Path = ROOT, folder: Path = FOLDER) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{LANGUAGES.get(lang, lang)}-review.xlsx"
    old = {}
    if path.exists():                                               # keep corrections a reader already typed
        ws = load_workbook(path).active
        old = {r[6]: str(r[5]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if len(r) > 6 and r[5] and str(r[5]).strip()}
    wb = Workbook(); ws = wb.active; ws.title = f"{LANGUAGES.get(lang, lang)} review"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows(lang, root):
        ws.append(list(r[:5]) + [old.get(r[6], "")] + [r[6]])
    for col, w in zip("ABCDEFG", (12, 26, 9, 55, 55, 55, 20)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[5].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "A2"; ws.column_dimensions["G"].hidden = True
    wb.save(path)
    return path


def apply(lang: str, root: Path = ROOT, folder: Path = FOLDER) -> dict:
    path = folder / f"{LANGUAGES.get(lang, lang)}-review.xlsx"
    ws = load_workbook(path).active
    changed = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[5] or not str(r[5]).strip():
            continue
        stem, idx, kind = r[6].split("|")
        f = root / "content" / "translations" / lang / f"{stem}.json"
        tr = json.loads(f.read_text())
        if kind == "vo":
            tr["scenes"][int(idx)]["vo"] = str(r[5]).strip()
        elif kind == "card":
            tr["scenes"][int(idx)]["card"] = str(r[5]).strip()
        else:
            tr[kind] = str(r[5]).strip()
        tr["reviewed"] = True; tr.pop("machine_translated", None)
        f.write_text(json.dumps(tr, indent=1, ensure_ascii=False) + "\n"); changed[stem] = changed.get(stem, 0) + 1
    return changed


def main(argv=None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2 or argv[0] not in ("export", "apply"):
        print(__doc__); return 2
    if argv[0] == "export":
        print("wrote", export(argv[1]))
    else:
        c = apply(argv[1]); print(f"applied corrections to {len(c)} scripts:", c)
    return 0


if __name__ == "__main__":
    sys.exit(main())
